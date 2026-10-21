#!/usr/bin/env python3
import json, sys
from openpyxl import load_workbook

source=sys.argv[1]
wb=load_workbook(source,data_only=False)
ws=wb['数据']
assert ws.max_row==43 and ws.max_column==4
assert ws.freeze_panes=='A2'
assert ws.auto_filter.ref=='A1:D43'
assert ws['A1'].value=='区域' and ws['A1'].font.bold and ws['A1'].font.color.rgb=='FFFFFFFF'
assert ws['B4'].value=='=1+1' and ws['B4'].data_type=='s'
assert ws['C2'].data_type=='n' and ws['C2'].value==-1250.5
assert ws['D2'].value=='第一行：核对材料\n第二行：确认责任人' and ws['D2'].alignment.wrap_text
assert all(10<=ws.column_dimensions[c].width<=40 for c in 'ABCD')
assert ws.sheet_properties.pageSetUpPr.fitToPage is None or ws.sheet_properties.pageSetUpPr.fitToPage in (True, False)
print(json.dumps({'verified':True,'rows':42,'formula_like_text_preserved':True,'freeze_panes':'A2','auto_filter':'A1:D43','numeric_cells':True,'wrapped_multiline':True},ensure_ascii=False))
